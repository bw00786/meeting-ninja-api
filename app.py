from flask import Flask, request, jsonify, send_file
import requests
import os
import re
import docx
from docx import Document
from flask_cors import CORS
from fpdf import FPDF
import logging

# Initialize Flask and enable CORS
app = Flask(__name__)

CORS(app)  # This enables CORS for all routes

# Load environment variables
url = os.getenv('LLM_URL')
backup_url = os.getenv('LLM_URL_BACKUP')
ssnc_api_key = os.getenv('SSC_CLOUD_API_KEY')
model = os.getenv('MODEL_NAME')
upload_folder = os.getenv('UPLOAD_FOLDER')  # Set this to a temporary folder
output_folder = os.getenv('OUTPUT_FOLDER')  # Set this to where you want the PDF files to be saved
http_port = os.getenv('HTTP_PORT')

app.config['UPLOAD_FOLDER'] = upload_folder
app.config['OUTPUT_FOLDER'] = output_folder

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Function to extract text from DOCX
def extract_text_from_docx(docx_path):
    """Extract text from a DOCX file."""
    try:
        doc = Document(docx_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    except Exception as e:
        logging.error(f"Error extracting text from DOCX: {e}")
        raise

# Function to analyze the meeting transcript using Llama 3.1, with fallback URL
def analyze_meeting_transcript(transcript):
    """Analyze the meeting transcript using the Llama 3.1 model, with a backup URL if the primary fails."""
    
    payload = {
        "model": model,
        "max_toke"
        "messages": [
            {"role": "system", "content": "You are an assistant that helps analyze meeting transcripts."},
            {"role": "user", "content": f"Analyze the following meeting transcript:\n\n{transcript}\n\nProvide very detailed meeting minutes with details like the date of the meeting, the speakers, what were the highlights of what the speakers said, the category, any conclusions, next steps, and action items. Make sure you highlight Attendees, Categories, conclusions and Meeting Summary in bold. The headings should be supported by pdf format, place the bold text in Markdown format."}
        ]
    }
    headers = {
        "Content-Type": "application/json",
        "Authorization": ssnc_api_key,
    }

    def attempt_request(url_to_try):
        """Helper function to attempt a request to the given LLM URL."""
        response = requests.post(url_to_try, json=payload, headers=headers)
        response.raise_for_status()  # Raise an error for bad responses
        result = response.json()
        # Check for the structure and existence of 'choices'
        if 'choices' in result and len(result['choices']) > 0:
            return result['choices'][0].get('message', {}).get('content', 'No response text found')
        else:
            error_message = result.get('choices', [{'message': {'content': 'Unknown error'}}])[0]['message']['content']
            logger.error(f"LLM returned an error: {error_message}")
            return f"LLM error: {error_message}"

    try:
        # Attempt to communicate with the primary LLM URL
        return attempt_request(url)
    except requests.exceptions.RequestException as e:
        logging.error(f"Error communicating with primary LLM: {e}")
        if backup_url:
            try:
                logging.info("Attempting to communicate with the backup LLM URL.")
                return attempt_request(backup_url)
            except requests.exceptions.RequestException as backup_error:
                logging.error(f"Error communicating with backup LLM: {backup_error}")
                raise Exception(f"Both primary and backup LLM URLs failed. Primary error: {e}, Backup error: {backup_error}")
        else:
            raise Exception(f"Primary LLM URL failed and no backup URL configured. Error: {e}")

# Function to write the meeting minutes to a PDF
def write_minutes_to_pdf(meeting_minutes, output_pdf_path):
    """Write the meeting minutes to a PDF file."""
    try:
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        # Set font for the body text
        pdf.set_font("Arial", size=12)

        # Add a title
        pdf.set_font("Arial", style='B', size=16)
        pdf.cell(200, 10, txt="Meeting Minutes", ln=True, align="C")
        pdf.ln(10)

        # Add the meeting minutes content
        pdf.set_font("Arial", size=14)  # Set the font size to 14 for the content
        for line in meeting_minutes.split('\n'):
            # Detect and process bold text in markdown-like format (**bold**)
            parts = re.split(r'(\*\*.*?\*\*)', line)  # Split by **bold** patterns
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    # If part is bold (in markdown-like syntax), make it bold in PDF
                    pdf.set_font("Arial", style='B', size=14)
                    pdf.multi_cell(0, 10, part[2:-2])  # Remove ** from the text
                else:
                    # Normal text
                    pdf.set_font("Arial", size=14)
                    pdf.multi_cell(0, 10, part)

        # Output the PDF to a file
        pdf.output(output_pdf_path)
    except Exception as e:
        logging.error(f"Error writing PDF: {e}")
        raise    

# API endpoint to generate meeting minutes from uploaded DOCX file
@app.route('/generate_minutes', methods=['POST'])
def generate_meeting_minutes():
    try:
        # Check if the request contains the file
        if 'docx_file' not in request.files:
            logging.warning("No DOCX file uploaded.")
            return jsonify({"error": "No DOCX file uploaded."}), 400

        file = request.files['docx_file']
        if file.filename == '':
            logging.warning("No file selected.")
            return jsonify({"error": "No file selected."}), 400

        if file and file.filename.endswith('.docx'):
            # Save the uploaded file
            docx_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
            file.save(docx_path)
            logging.info(f"File {file.filename} uploaded successfully.")

            # Define the output PDF path
            output_pdf_path = os.path.join(app.config['OUTPUT_FOLDER'], f"{os.path.splitext(file.filename)[0]}.pdf")

            # Process the file and generate the PDF
            transcript = extract_text_from_docx(docx_path)
            meeting_minutes = analyze_meeting_transcript(transcript)
            write_minutes_to_pdf(meeting_minutes, output_pdf_path)

            logging.info(f"Meeting minutes generated and saved to {output_pdf_path}.")
            return jsonify({
                "message": "PDF generated successfully",
                "filename": os.path.basename(output_pdf_path),
                "fullPath": output_pdf_path
            })

        logging.warning("Invalid file type. Only DOCX files are allowed.")
        return jsonify({"error": "Invalid file type. Please upload a DOCX file."}), 400

    except Exception as e:
        logging.error(f"Error in generate_meeting_minutes: {e}")
        return jsonify({"error": str(e)}), 500
    
# API endpoint to download the generated PDF
@app.route('/download_pdf/<filename>', methods=['GET'])
def download_pdf(filename):
    try:
        file_path = os.path.join(app.config['OUTPUT_FOLDER'], filename)
        app.logger.info(f"Attempting to serve file: {file_path}")
        
        if not os.path.exists(file_path):
            app.logger.error(f"File not found: {file_path}")
            raise NotFound(f"File not found: {filename}")
        
        return send_file(file_path,
                         mimetype='application/pdf',
                         as_attachment=True,
                         download_name=filename)
    except NotFound as e:
        app.logger.error(f"NotFound error in download_pdf: {str(e)}")
        return jsonify({"error": str(e)}), 404
    except Exception as e:
        app.logger.error(f"Unexpected error in download_pdf: {str(e)}")
        return jsonify({"error": f"An unexpected error occurred: {str(e)}"}), 500    

if __name__ == "__main__":
    app.run(host='0.0.0.0', port=http_port)
